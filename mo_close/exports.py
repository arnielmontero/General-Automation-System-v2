"""Excel (openpyxl) and PDF (ReportLab) reports for MO's To Be Close."""

import io

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.models import AuditLog

from .models import SYSTEM
from .services import progress_counts

LOGO = settings.BASE_DIR / "static" / "img" / "sawo-logo.png"

ROW_COLUMNS = [
    ("SO#", 10),
    ("FG Code", 20),
    ("FG", 38),
    ("Qty", 7),
    ("Type", 7),
    ("Reference", 17),
    ("Odoo State", 22),
    ("Marked Done", 8),
    ("Marked Done Date", 17),
    ("Marked By", 18),
    ("System Check", 16),
    ("Checked At", 17),
    ("Final Confirmed", 9),
    ("Confirmed By", 18),
    ("Confirmed Date", 17),
    ("Comment", 36),
    ("Source Fix", 20),
]

STAGE_FILL = {
    "confirmed": "E8F7EE",  # green
    "waiting": "FFF8DB",  # yellow
    "flag": "FDECEC",  # red
}


def _dt(value):
    return timezone.localtime(value).strftime("%d-%b-%Y %H:%M") if value else ""


def _num(value):
    return float(value) if value is not None else None


def _who(user):
    return user.display_name if user else ""


def _stage(row):
    if row.final_confirmed:
        return "confirmed"
    if row.user_done and row.system_status != "confirmed":
        return "flag"
    if row.user_done:
        return "waiting"
    return "open"


def _so_label(pl):
    """'S04743 (+ S04743-1, S08869)': the parent SO with its child SOs."""
    children = [n.so_number for n in pl.child_sos]
    return f"{pl.so_number} (+ {', '.join(children)})" if children else pl.so_number


def record_rows(pl, only_open=False):
    """Flat rows: each MO followed by its transfers, by SO as on the record page (an MO serving the parent and an
    additional packing list is listed under both, each with its own qty)."""
    out = []
    mos = pl.mos.select_related("packing_list", "sales_order", "user_done_by", "final_confirmed_by").prefetch_related(
        "transfers__user_done_by", "transfers__final_confirmed_by"
    )
    for mo, node, qty in pl.mo_placements(list(mos)):
        so = node.so_number if node else mo.so_number
        for row in [mo, *mo.transfers.all()]:
            if only_open and row.final_confirmed:
                continue
            is_mo = row.KIND == "mo"
            out.append(
                {
                    "row": row,
                    "stage": _stage(row),
                    "values": [
                        so,
                        mo.fg_code,
                        (mo.fg_description if mo.on_packing_list else f"{mo.fg_description} (component — not on packing list)")
                        if is_mo else "",
                        _num(qty) if is_mo and mo.on_packing_list else None,
                        "MO" if is_mo else (row.type_code or "Transfer"),
                        row.reference,
                        row.state_label,
                        "Yes" if row.user_done else "No",
                        _dt(row.user_done_at),
                        _who(row.user_done_by),
                        row.get_system_status_display() if row.system_status else "",
                        _dt(row.system_checked_at),
                        "Yes" if row.final_confirmed else "No",
                        _who(row.final_confirmed_by),
                        _dt(row.final_confirmed_at),
                        row.final_comment,
                        f"{mo.source_fix} → {mo.source_so} ({mo.source_fix_hint})" if is_mo and mo.source_fix else "",
                    ],
                }
            )
    return out


# --- Excel ------------------------------------------------------------------------------------

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1D2330")
THIN = Side(style="thin", color="D9DCE1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet_table(ws, start_row, columns, rows, fills=None):
    for c, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=title)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, values in enumerate(rows, start=start_row + 1):
        fill = fills[r - start_row - 1] if fills else None
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border, cell.alignment = BORDER, WRAP
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{start_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(color="8A91A0")


def _info_pairs(pl, counts):
    return [
        ("SO", pl.so_number),
        ("Child SOs", ", ".join(n.path for n in pl.child_sos) or "—"),
        ("Packing lists", pl.documents.filter(saved=True).count()),
        ("Status", pl.get_status_display()),
        ("Sold To", pl.sold_to.replace("\n", ", ")),
        ("Tel. No.", pl.tel),
        ("Fax No.", pl.fax),
        ("ATTN", pl.attn),
        ("Date", pl.doc_date_text),
        ("Invoice No.", pl.invoice_no),
        ("Container No.", pl.container_no),
        ("Seal No.", pl.seal_no),
        ("Destination", pl.destination),
        ("Total Qty (PDF)", _num(pl.total_qty)),
        ("Total Packages (PDF)", _num(pl.total_packages)),
        ("Created", f"{_dt(pl.created_at)} by {_who(pl.created_by)}"),
        ("Generated from Odoo", f"{_dt(pl.generated_at)} by {_who(pl.generated_by)}" if pl.generated_at else "Not yet"),
        ("Last refresh", _dt(pl.refreshed_at)),
        ("Rows (MO + transfers)", counts["total"]),
        ("Marked done", counts["marked"]),
        ("Done in Odoo", counts["system"]),
        ("Final confirmed", counts["final"]),
        ("Progress", f"{counts['percent']}%"),
    ]


def record_excel(pl):
    counts = progress_counts([pl])[pl.pk]
    wb = Workbook()

    ws = wb.active
    ws.title = "Record"
    _title(ws, f"MO's To Be Close — {pl.so_number}", f"Exported {_dt(timezone.now())}")
    for r, (label, value) in enumerate(_info_pairs(pl, counts), start=4):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    ws = wb.create_sheet("MO Progress")
    rows = record_rows(pl)
    _title(ws, f"MO's and transfers — {pl.so_number}", "Green = confirmed done · Yellow = waiting final confirmation · Red = not confirmed by Odoo")
    _sheet_table(ws, 4, ROW_COLUMNS, [r["values"] for r in rows], [STAGE_FILL.get(r["stage"]) for r in rows])

    docs = [
        (node, doc)
        for node in pl.so_tree()
        for doc in pl.documents.filter(sales_order=node, saved=True).prefetch_related("lines")
    ]
    ws = wb.create_sheet("Packing List")
    _title(ws, f"Packing list — {pl.so_number}", "By SO (parent first, then child SOs) · grouped by External Code")
    _sheet_table(
        ws,
        4,
        [("SO", 14), ("Parent SO", 14), ("Invoice No", 12), ("External Code", 22), ("Description", 48), ("Qty", 8),
         ("PDF rows", 9), ("PDF", 30)],
        [
            [node.so_number, node.parent.so_number if node.parent_id else "", doc.invoice_no, g["code"],
             g["description"], _num(g["qty"]), g["rows"], doc.pdf_name]
            for node, doc in docs
            for g in doc.grouped_lines()
        ],
    )

    ws = wb.create_sheet("PDF Rows")
    _title(ws, f"PDF rows — {pl.so_number}", f"{len(docs)} packing list PDF(s)")
    _sheet_table(
        ws,
        4,
        [("SO", 14), ("PDF", 30), ("#", 5), ("External Code", 22), ("Description", 48), ("UM", 8), ("Packages", 10),
         ("Net Weight", 11), ("Gross Weight", 12), ("CBM", 9), ("Qty", 8)],
        [
            [node.so_number, doc.pdf_name, l.line_no, l.external_code, l.description, l.um, _num(l.packages),
             _num(l.net_weight), _num(l.gross_weight), _num(l.cbm), _num(l.qty)]
            for node, doc in docs
            for l in doc.lines.all()
        ],
    )

    ws = wb.create_sheet("Logs")
    _title(ws, f"Logs — {pl.so_number}", "")
    logs = AuditLog.objects.filter(system=SYSTEM, record_id=pl.pk)
    _sheet_table(
        ws,
        4,
        [("Date", 18), ("User", 20), ("Section", 22), ("Action", 28), ("Item", 18), ("Old", 20), ("New", 30), ("Comment", 40)],
        [
            [_dt(e.created_at), e.user_label, e.get_section_display(), e.action, e.target, e.old_value, e.new_value, e.comment]
            for e in logs
        ],
    )
    return _save(wb)


def open_items_excel(records):
    counts = progress_counts(records)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _title(ws, "MO's To Be Close — open records", f"Exported {_dt(timezone.now())}")
    _sheet_table(
        ws,
        4,
        [("SO", 12), ("Sold To", 30), ("Container No.", 18), ("Rows", 8), ("Marked done", 12),
         ("Done in Odoo", 12), ("Final confirmed", 14), ("Progress %", 11)],
        [
            [_so_label(pl), pl.sold_to_name, pl.container_no, counts[pl.pk]["total"], counts[pl.pk]["marked"],
             counts[pl.pk]["system"], counts[pl.pk]["final"], counts[pl.pk]["percent"]]
            for pl in records
        ],
    )
    ws = wb.create_sheet("Not Done Yet")
    rows = [r for pl in records for r in record_rows(pl, only_open=True)]
    _title(ws, "Rows not yet confirmed done", "Yellow = waiting final confirmation · Red = not confirmed by Odoo")
    _sheet_table(ws, 4, ROW_COLUMNS, [r["values"] for r in rows], [STAGE_FILL.get(r["stage"]) for r in rows])
    return _save(wb)


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------------------------

STYLE = ParagraphStyle("cell", fontName="Helvetica", fontSize=7, leading=8.5)
STYLE_B = ParagraphStyle("cellb", parent=STYLE, fontName="Helvetica-Bold")
STYLE_TITLE = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=14, leading=17)
STYLE_SUB = ParagraphStyle("sub", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#6B7280"))
PDF_FILLS = {k: colors.HexColor("#" + v) for k, v in STAGE_FILL.items()}

PDF_COLUMNS = [  # (title, width mm, index in ROW_COLUMNS values)
    ("FG Code", 30, 1), ("FG", 50, 2), ("Qty", 9, 3), ("Type", 10, 4), ("Reference", 26, 5), ("Odoo State", 26, 6),
    ("Marked Done", 30, None), ("System Check", 24, 10), ("Final Confirmation", 36, None), ("Comment", 36, 15),
]


def _p(text, style=STYLE):
    return Paragraph(("" if text is None else str(text)).replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>"), style)


def _pdf_row(r):
    v = r["values"]
    cells = []
    for _, _, idx in PDF_COLUMNS:
        if idx is not None:
            value = v[idx]
            if idx == 3 and value is not None:
                value = f"{value:g}"
            cells.append(_p(value, STYLE_B if idx == 5 and v[4] == "MO" else STYLE))
    marked = f"Yes · {v[8]}\n{v[9]}" if v[7] == "Yes" else "No"
    final = f"Yes · {v[14]}\n{v[13]}" if v[12] == "Yes" else "No"
    cells.insert(6, _p(marked))
    cells.insert(8, _p(final))
    return cells


def _header_block(title, subtitle):
    logo = Image(str(LOGO), width=26 * mm, height=16 * mm) if LOGO.exists() else ""
    t = Table([[logo, [Paragraph(title, STYLE_TITLE), Paragraph(subtitle, STYLE_SUB)]]], colWidths=[32 * mm, None])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    return t


def _rows_table(rows):
    data = [[_p(title, STYLE_B) for title, _, _ in PDF_COLUMNS]] + [_pdf_row(r) for r in rows]
    table = Table(data, colWidths=[w * mm for _, w, _ in PDF_COLUMNS], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F2F4")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]
    for i, r in enumerate(rows, start=1):
        fill = PDF_FILLS.get(r["stage"])
        if fill:
            style.append(("BACKGROUND", (0, i), (-1, i), fill))
    table.setStyle(TableStyle(style))
    return table


def _build_pdf(story):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
        title="MO's To Be Close", author="SAWO General Automation System",
    )

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#8A91A0"))
        canvas.drawString(10 * mm, 6 * mm, f"MO's To Be Close · exported {_dt(timezone.now())}")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _info_table(pl, counts):
    pairs = _info_pairs(pl, counts)
    left, right = pairs[:11], pairs[11:]
    data = []
    for i in range(max(len(left), len(right))):
        row = []
        for side in (left, right):
            if i < len(side):
                value = side[i][1]
                if isinstance(value, float):
                    value = f"{value:g}"
                row += [_p(side[i][0], STYLE_B), _p(value)]
            else:
                row += ["", ""]
        data.append(row)
    t = Table(data, colWidths=[30 * mm, 100 * mm, 36 * mm, 100 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
    return t


def record_pdf(pl):
    counts = progress_counts([pl])[pl.pk]
    story = [
        _header_block(f"MO's To Be Close — {pl.so_number}", f"{pl.sold_to_name} · {pl.get_status_display()} · {counts['percent']}% confirmed"),
        Spacer(1, 4 * mm),
        _info_table(pl, counts),
        Spacer(1, 5 * mm),
    ]
    rows = record_rows(pl)
    if rows:
        story.append(_rows_table(rows))
    else:
        story.append(_p("Not generated from Odoo yet.", STYLE_SUB))
    return _build_pdf(story)


def open_items_pdf(records):
    counts = progress_counts(records)
    story = [_header_block("MO's To Be Close — open items", f"{len(records)} record(s) in progress · rows not yet confirmed done"), Spacer(1, 4 * mm)]
    summary = [[_p(t, STYLE_B) for t in ("SO", "Sold To", "Container No.", "Rows", "Marked done", "Done in Odoo", "Final confirmed", "Progress")]]
    for pl in records:
        c = counts[pl.pk]
        summary.append([_p(x) for x in (_so_label(pl), pl.sold_to_name, pl.container_no, c["total"], c["marked"], c["system"], c["final"], f"{c['percent']}%")])
    t = Table(summary, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F2F4")),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE1"))]))
    story.append(t)
    for pl in records:
        rows = record_rows(pl, only_open=True)
        if not rows:
            continue
        story += [PageBreak(), _header_block(f"{_so_label(pl)} — {pl.sold_to_name}", f"{len(rows)} row(s) not yet confirmed done"),
                  Spacer(1, 4 * mm), _rows_table(rows)]
    return _build_pdf(story)
