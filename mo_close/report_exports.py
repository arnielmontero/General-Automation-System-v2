"""Reports page downloads: Excel with native charts (openpyxl) and PDF with charts (ReportLab graphics)."""

from math import ceil

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.widgets.markers import makeMarker
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, Spacer, Table, TableStyle

from .exports import (
    BORDER, HEAD_FILL, HEAD_FONT, STYLE_B, STYLE_SUB, STYLE_TITLE,
    _build_pdf, _dt, _header_block, _p, _save, _sheet_table, _so_label, _title,
)
from .reports import RECORD_SERIES

PERIOD_COLUMNS = [("Period", 16), ("SO records", 11), ("In progress", 11), ("Not generated", 13), ("Closed", 9),
                  ("Rows", 8), ("Final confirmed (rows)", 14), ("Confirmations in period", 14)]
SO_COLUMNS = [("SO", 11), ("Sold To", 28), ("Date", 13), ("Status", 13), ("Rows", 8), ("Marked done", 11),
              ("Done in Odoo", 11), ("Final confirmed", 12), ("% confirmed", 11), ("Days open", 10)]


def _period_values(p):
    return [p["label"], p["records"], p["generated"], p["saved"], p["closed"], p["rows"], p["final"], p["confirmations"]]


def _so_values(s):
    pl = s["pl"]
    return [_so_label(pl), pl.sold_to_name, s["date"].strftime("%d-%b-%Y") if s["date"] else "",
            pl.get_status_display(), s["total"], s["marked"], s["system"], s["final"], s["percent"], s["days_open"]]


# --- Excel ------------------------------------------------------------------------------------


def _xl(hex_color):
    return hex_color.lstrip("#").upper()


def _xl_pie(ws, title, items, top_row, col, anchor):
    """Small table (the chart's table view) plus a pie chart next to it."""
    ws.cell(row=top_row, column=col, value=title).font = Font(bold=True)
    head = top_row + 1
    for c, text in enumerate(("Item", "Count"), start=col):
        cell = ws.cell(row=head, column=c, value=text)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
    for r, item in enumerate(items, start=head + 1):
        ws.cell(row=r, column=col, value=item["label"]).border = BORDER
        ws.cell(row=r, column=col + 1, value=item["value"]).border = BORDER
    if not sum(i["value"] for i in items):
        return
    pie = PieChart()
    pie.title = title
    pie.height, pie.width = 7, 12
    pie.add_data(Reference(ws, min_col=col + 1, min_row=head, max_row=head + len(items)), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=col, min_row=head + 1, max_row=head + len(items)))
    for idx, item in enumerate(items):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = _xl(item["color"])
        point.graphicalProperties.line.solidFill = "FFFFFF"
        pie.series[0].dPt.append(point)
    ws.add_chart(pie, anchor)


def report_excel(data):
    f, k = data["filters"], data["kpis"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    _title(ws, "MO's To Be Close — Report", f"{f.description} · exported {_dt(timezone.now())}")
    summary = [
        ("SO records", k["records"]), ("In progress", k["in_progress"]), ("Not generated yet", k["not_generated"]),
        ("Closed", k["closed"]), ("MOs", k["mos"]), ("Transfers", k["transfers"]), ("Rows (MO + transfers)", k["rows"]),
        ("Marked done by users", k["marked"]), ("Done in Odoo", k["odoo_done"]), ("Final confirmed", k["final"]),
        ("% final confirmed", f'{k["percent"]}%'), ("Avg. days to close", k["avg_days_to_close"] or ""),
    ]
    for r, (label, value) in enumerate(summary, start=4):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
    for col, width in (("A", 26), ("B", 12), ("D", 28), ("E", 10)):
        ws.column_dimensions[col].width = width
    _xl_pie(ws, "Row progress", data["stages"], 4, 4, "G4")
    _xl_pie(ws, "MO states", data["mo_states"], 11, 4, "G19")
    _xl_pie(ws, "Transfer states", data["transfer_states"], 19, 4, "G34")

    ws = wb.create_sheet("By period")
    _title(ws, f"By {f.group}", f.description)
    periods = data["periods"]
    _sheet_table(ws, 4, PERIOD_COLUMNS, [_period_values(p) for p in periods])
    if periods:
        last = 4 + len(periods)
        bar = BarChart()
        bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
        bar.title = f"SO records per {f.group}"
        bar.height, bar.width = 8, 18
        bar.add_data(Reference(ws, min_col=3, max_col=5, min_row=4, max_row=last), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
        for series, (_, _, color) in zip(bar.series, RECORD_SERIES):
            series.graphicalProperties.solidFill = _xl(color)
            series.graphicalProperties.line.solidFill = "FFFFFF"
        ws.add_chart(bar, "J4")

        line = LineChart()
        line.title = f"Final confirmations per {f.group}"
        line.height, line.width = 8, 18
        line.legend = None
        line.add_data(Reference(ws, min_col=8, min_row=4, max_row=last), titles_from_data=True)
        line.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
        line.series[0].graphicalProperties.line.solidFill = "2A78D6"
        line.series[0].graphicalProperties.line.width = 22000
        ws.add_chart(line, "J22")

    ws = wb.create_sheet("By SO")
    _title(ws, "By SO", f.description)
    per_so = data["per_so"]
    _sheet_table(ws, 4, SO_COLUMNS, [_so_values(s) for s in per_so])
    if per_so:
        hbar = BarChart()
        hbar.type = "bar"
        hbar.title = "% confirmed by SO"
        hbar.legend = None
        hbar.height, hbar.width = max(7, 0.55 * len(per_so) + 3), 16
        hbar.y_axis.scaling.min, hbar.y_axis.scaling.max = 0, 100
        hbar.add_data(Reference(ws, min_col=9, min_row=4, max_row=4 + len(per_so)), titles_from_data=True)
        hbar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(per_so)))
        hbar.series[0].graphicalProperties.solidFill = "2A78D6"
        ws.add_chart(hbar, "M4")
    return _save(wb)


# --- PDF --------------------------------------------------------------------------------------

INK = colors.HexColor("#1D2330")
MUTED = colors.HexColor("#898781")
GRID = colors.HexColor("#E1E0D9")
AXIS = colors.HexColor("#C3C2B7")
BLUE = colors.HexColor("#2a78d6")


def _thin(labels, max_n=10):
    """Show at most ~10 axis labels so they never collide."""
    if len(labels) <= max_n:
        return labels
    step = ceil(len(labels) / max_n)
    return [label if i % step == 0 else "" for i, label in enumerate(labels)]


def _drawing(w, h, title):
    d = Drawing(w, h)
    d.add(String(0, h - 10, title, fontName="Helvetica-Bold", fontSize=9, fillColor=INK))
    return d


def _no_data(d, w, h):
    d.add(String(w / 2, h / 2, "No data for these filters", textAnchor="middle",
                 fontName="Helvetica", fontSize=8, fillColor=MUTED))
    return d


def _value_axis(axis, top):
    step = max(1, ceil(top / 5))
    axis.valueMin, axis.valueMax, axis.valueStep = 0, step * max(1, ceil(top / step)), step
    axis.visibleGrid, axis.gridStrokeColor, axis.gridStrokeWidth = True, GRID, 0.5
    axis.visibleAxis = axis.visibleTicks = False
    axis.labels.fontName, axis.labels.fontSize, axis.labels.fillColor = "Helvetica", 6.5, MUTED


def _category_axis(axis, labels):
    axis.categoryNames = _thin(labels)
    axis.strokeColor, axis.strokeWidth = AXIS, 0.5
    axis.visibleTicks = False
    axis.labels.fontName, axis.labels.fontSize, axis.labels.fillColor = "Helvetica", 6.5, MUTED
    axis.labels.boxAnchor = "n"


def _legend(x, y, pairs, columns=1):
    lg = Legend()
    lg.x, lg.y, lg.boxAnchor = x, y, "nw"
    lg.fontName, lg.fontSize = "Helvetica", 7
    lg.colorNamePairs = pairs
    lg.alignment = "right"
    lg.dx, lg.dy, lg.dxTextSpace, lg.deltay, lg.deltax = 6, 6, 4, 10, 80
    lg.strokeColor = None
    lg.columnMaximum = ceil(len(pairs) / columns)
    return lg


def _periods_chart(data, w, h):
    p, f = data["chart"]["periods"], data["filters"]
    d = _drawing(w, h, f"SO records per {f.group}")
    totals = [sum(v) for v in zip(*[s["data"] for s in p["series"]])] if p["labels"] else []
    if not totals or not max(totals):
        return _no_data(d, w, h)
    bc = VerticalBarChart()
    bc.x, bc.y, bc.width, bc.height = 28, 22, w - 36, h - 62
    bc.data = [s["data"] for s in p["series"]]
    bc.categoryAxis.style = "stacked"
    bc.groupSpacing = 4
    for i, s in enumerate(p["series"]):
        bc.bars[i].fillColor = colors.HexColor(s["color"])
        bc.bars[i].strokeColor, bc.bars[i].strokeWidth = colors.white, 0.6
    _value_axis(bc.valueAxis, max(totals))
    _category_axis(bc.categoryAxis, p["labels"])
    d.add(bc)
    d.add(_legend(28, h - 18, [(colors.HexColor(s["color"]), s["label"]) for s in p["series"]], columns=3))
    return d


def _confirmations_chart(data, w, h):
    p, f = data["chart"]["periods"], data["filters"]
    d = _drawing(w, h, f"Final confirmations per {f.group}")
    if not p["labels"] or not max(p["confirmations"] or [0]):
        return _no_data(d, w, h)
    lc = HorizontalLineChart()
    lc.x, lc.y, lc.width, lc.height = 28, 22, w - 36, h - 50
    lc.data = [p["confirmations"]]
    lc.joinedLines = 1
    lc.lines[0].strokeColor, lc.lines[0].strokeWidth = BLUE, 1.5
    lc.lines[0].symbol = makeMarker("FilledCircle", size=4, fillColor=BLUE, strokeColor=colors.white)
    _value_axis(lc.valueAxis, max(p["confirmations"]))
    _category_axis(lc.categoryAxis, p["labels"])
    d.add(lc)
    return d


def _pie_chart(title, items, w, h):
    d = _drawing(w, h, title)
    total = sum(i["value"] for i in items)
    if not total:
        return _no_data(d, w, h)
    shown = [i for i in items if i["value"]]
    size = min(h - 30, w * 0.42)
    pie = Pie()
    pie.x, pie.y, pie.width, pie.height = 4, (h - 18 - size) / 2, size, size
    pie.data = [i["value"] for i in shown]
    pie.labels = None
    pie.slices.strokeColor, pie.slices.strokeWidth = colors.white, 1
    for n, i in enumerate(shown):
        pie.slices[n].fillColor = colors.HexColor(i["color"])
    d.add(pie)
    pairs = [(colors.HexColor(i["color"]), f'{i["label"]}  {i["value"]} ({round(100 * i["value"] / total)}%)') for i in items]
    d.add(_legend(size + 18, h - 26, pairs))
    return d


def _by_so_chart(data, w, h_max):
    s = data["chart"]["by_so"]
    labels, values = s["labels"][:25], s["data"][:25]
    h = min(h_max, max(90, 13 * len(labels) + 50))
    d = _drawing(w, h, "% confirmed by SO" + (" (25 lowest)" if len(s["labels"]) > 25 else ""))
    if not labels:
        return _no_data(d, w, h)
    bc = HorizontalBarChart()
    bc.x, bc.y, bc.width, bc.height = 48, 18, w - 60, h - 40
    bc.data = [values]
    bc.bars[0].fillColor, bc.bars[0].strokeColor = BLUE, None
    bc.valueAxis.valueMin, bc.valueAxis.valueMax, bc.valueAxis.valueStep = 0, 100, 20
    bc.valueAxis.visibleGrid, bc.valueAxis.gridStrokeColor, bc.valueAxis.gridStrokeWidth = True, GRID, 0.5
    bc.valueAxis.visibleAxis = bc.valueAxis.visibleTicks = False
    bc.valueAxis.labelTextFormat = "%d%%"
    bc.valueAxis.labels.fontName, bc.valueAxis.labels.fontSize, bc.valueAxis.labels.fillColor = "Helvetica", 6.5, MUTED
    bc.categoryAxis.categoryNames = labels
    bc.categoryAxis.strokeColor, bc.categoryAxis.strokeWidth = AXIS, 0.5
    bc.categoryAxis.visibleTicks = False
    bc.categoryAxis.labels.fontName, bc.categoryAxis.labels.fontSize, bc.categoryAxis.labels.fillColor = "Helvetica", 6.5, INK
    bc.categoryAxis.labels.boxAnchor = "e"
    d.add(bc)
    return d


def _kpi_table(k):
    tiles = [
        ("SO records", k["records"], f'{k["in_progress"]} in progress · {k["closed"]} closed'),
        ("Rows (MO + transfers)", k["rows"], f'{k["mos"]} MOs · {k["transfers"]} transfers'),
        ("Marked done", k["marked"], ""),
        ("Done in Odoo", k["odoo_done"], ""),
        ("Final confirmed", k["final"], f'{k["percent"]}% of rows'),
        ("Avg. days to close", k["avg_days_to_close"] if k["avg_days_to_close"] is not None else "—", "closed SOs"),
    ]
    big = ParagraphStyle("kpi", fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=INK)
    cells = [[[_p(label, STYLE_SUB), Paragraph(str(value), big), _p(note, STYLE_SUB)] for label, value, note in tiles]]
    t = Table(cells, colWidths=[46 * mm] * 6)
    t.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return t


def _simple_table(columns, rows, widths):
    body = [[_p(c, STYLE_B) for c, _ in columns]] + [[_p(v) for v in r] for r in rows]
    t = Table(body, colWidths=[w * mm for w in widths], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F2F4")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return t


def _side_by_side(left, right, half_w):
    t = Table([[left, right]], colWidths=[half_w + 4 * mm, half_w + 3 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    return t


def report_pdf(data):
    f = data["filters"]
    half = 135 * mm
    story = [
        _header_block("MO's To Be Close — Report", f.description),
        Spacer(1, 4 * mm),
        _kpi_table(data["kpis"]),
        Spacer(1, 6 * mm),
        _side_by_side(_periods_chart(data, half, 70 * mm), _pie_chart("Row progress", data["stages"], half, 70 * mm), half),
        Spacer(1, 5 * mm),
        _side_by_side(_pie_chart("MO states", data["mo_states"], half, 52 * mm),
                      _pie_chart("Transfer states", data["transfer_states"], half, 52 * mm), half),
        PageBreak(),
        _side_by_side(_confirmations_chart(data, half, 70 * mm), _by_so_chart(data, half, 170 * mm), half),
        PageBreak(),
        Paragraph(f"By {f.group}", STYLE_TITLE),
        Spacer(1, 3 * mm),
        _simple_table(PERIOD_COLUMNS, [_period_values(p) for p in data["periods"]], [40, 28, 28, 30, 26, 24, 40, 40])
        if data["periods"] else _p("No data for these filters.", STYLE_SUB),
        PageBreak(),
        Paragraph("By SO", STYLE_TITLE),
        Spacer(1, 3 * mm),
        _simple_table(SO_COLUMNS, [_so_values(s) for s in data["per_so"]], [22, 60, 24, 24, 18, 22, 24, 26, 22, 20])
        if data["per_so"] else _p("No SO records for these filters.", STYLE_SUB),
    ]
    return _build_pdf(story)
